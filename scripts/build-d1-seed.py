#!/usr/bin/env python3
"""Build the D1 seed from the canonical Google Sheets workbook export."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


AFFINITY_COLUMNS = [
    "ta_mon_thu_00_05", "ta_mon_thu_06_10", "ta_mon_thu_10_13", "ta_mon_thu_13_17",
    "ta_mon_thu_17_20", "ta_mon_thu_20_24", "ta_friday_00_05", "ta_friday_06_10",
    "ta_friday_10_13", "ta_friday_13_17", "ta_friday_17_20", "ta_friday_20_24",
    "ta_weekend_00_05", "ta_weekend_06_10", "ta_weekend_10_13", "ta_weekend_13_17",
    "ta_weekend_17_20", "ta_weekend_20_24",
]


def records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    return [dict(zip(headers, row)) for row in rows]


def text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return str(value).strip()


def number(value, default=0):
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return default


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + text(value).replace("'", "''") + "'"


def slugify(value):
    normal = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"(^-|-$)", "", re.sub(r"[^a-z0-9]+", "-", normal))[:160] or "place"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=False, data_only=True)
    master = records(workbook["POI_MASTER"])
    source = {text(row.get("poi_id")): row for row in records(workbook["POI_SOURCE_FULL"]) if text(row.get("poi_id"))}
    used_slugs = set()
    place_lines, affinity_lines, station_lines = [], [], []

    for row in master:
        place_id = text(row.get("poi_id"))
        name = text(row.get("name"))
        if not place_id or not name:
            continue
        src = source.get(place_id, {})
        base_slug = slugify(name)
        place_slug = base_slug
        counter = 2
        while place_slug in used_slugs:
            place_slug = f"{base_slug}-{counter}"
            counter += 1
        used_slugs.add(place_slug)
        opening_hours = {
            day: {"open": text(row.get(f"{day}_open")), "close": text(row.get(f"{day}_close"))}
            for day in ("mon", "tue", "wed", "thu", "fri", "sat", "sun")
            if text(row.get(f"{day}_open")) or text(row.get(f"{day}_close"))
        }
        lat, lon = float(row["lat"]), float(row["lon"])
        active = row.get("active") is True or text(row.get("active")).upper() == "TRUE"
        columns = [
            "id", "slug", "name", "category", "latitude", "longitude", "status", "description_en",
            "editorial_hook_en", "official_url", "guide_url", "google_maps_url", "price_text", "access_type",
            "access_notes_en", "opening_hours_json", "tourist_intensity", "crowd_scope", "visit_minutes",
            "best_time", "weather_fit", "visit_mode", "mood_quiet", "mood_unexpected", "mood_beautiful",
            "mood_weird", "mood_local", "mood_green", "mood_atmospheric", "mood_lively", "mood_reviewed",
            "mood_confidence", "mood_basis", "source_type", "source_ref", "source_row", "source_code_r1",
            "last_verified_at", "published_at", "created_at", "updated_at",
        ]
        values = [
            place_id, place_slug, name, text(row.get("category")), lat, lon,
            "published" if active else "draft", text(src.get("description")), text(row.get("editorial_hook")),
            text(row.get("official_url")), text(row.get("guide_url")), f"https://www.google.com/maps?q={lat},{lon}",
            text(src.get("Price")), text(row.get("access_type")) or "VARIABLE", text(row.get("access_notes")),
            json.dumps(opening_hours, ensure_ascii=False, separators=(",", ":")), number(row.get("tourist_intensity"), 30),
            text(row.get("crowd_scope")) or "VENUE", number(row.get("min_visit_minutes"), 30),
            text(row.get("best_time")) or "ANY", text(row.get("weather_fit")) or "ANY",
            text(row.get("visit_mode")) or "STOP", *[number(row.get(f"mood_{mood}"), 0) for mood in
            ("quiet", "unexpected", "beautiful", "weird", "local", "green", "atmospheric", "lively")],
            1 if row.get("mood_reviewed") is True else 0, text(row.get("mood_confidence")) or "LOW",
            text(row.get("mood_basis")), "google-sheet-migration", "London Advanced - Mood / POI_MASTER",
            number(src.get("source_row"), 0) or None, text(src.get("Code R1")), text(row.get("last_verified")) or None,
            "2026-09-25T00:00:00Z" if active else None, "2026-09-25T00:00:00Z", "2026-09-25T00:00:00Z",
        ]
        place_lines.append(f"INSERT INTO places({','.join(columns)}) VALUES ({','.join(sql(value) for value in values)});")

        affinity_values = [number(row.get(column), 50) for column in AFFINITY_COLUMNS]
        affinity_lines.append(
            "INSERT INTO place_time_affinity(place_id," + ",".join(column.removeprefix("ta_") for column in AFFINITY_COLUMNS)
            + ") VALUES (" + ",".join([sql(place_id), *map(str, affinity_values)]) + ");"
        )
        for position in (1, 2, 3):
            station_name = text(row.get(f"station{position}_name"))
            distance = row.get(f"station{position}_distance_m")
            if station_name and distance is not None:
                station_lines.append(
                    "INSERT INTO place_stations(place_id,position,station_id,station_name,distance_metres) VALUES ("
                    + ",".join([sql(place_id), str(position), sql(text(row.get(f"station{position}_id"))),
                                sql(station_name), str(number(distance))]) + ");"
                )

    if len(place_lines) != 881:
        raise RuntimeError(f"Expected 881 places, found {len(place_lines)}")
    output = ["PRAGMA foreign_keys=ON;", "BEGIN TRANSACTION;", *place_lines, *affinity_lines, *station_lines,
              "UPDATE app_meta SET value='1', updated_at=CURRENT_TIMESTAMP WHERE key='data_version';", "COMMIT;"]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(output) + "\n", encoding="utf-8")
    print(json.dumps({"places": len(place_lines), "affinity": len(affinity_lines), "stations": len(station_lines),
                      "output": str(args.output)}, indent=2))


if __name__ == "__main__":
    main()
